"""Render docs/env_v2/history/SYNTHETIC_ENVIRONMENT_HISTORY.md as a formatted .docx, and export the .pdf through Word.

Builds on tools.md_to_docx (headings, paragraphs, lists, pipe tables, inline formatting) and adds what the history
document needs: a title page, a real Word table of contents, embedded figures with their captions, repeated table
header rows, A4 page setup with page numbers. Content is copied verbatim; only formatting is
added.

Usage: python -m tools.history_docx [--md <input.md>] [--docx <out.docx>] [--pdf <out.pdf>] [--no-pdf]

The PDF is exported by Microsoft Word through COM (pywin32): the .docx is opened, every field and the table of
contents are updated, the .docx is saved again with the populated table of contents, and the PDF is written with
Word's own exporter. Word is closed afterwards.
"""
from __future__ import annotations

import argparse
import os
import re
import sys
import time

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from tools.md_to_docx import GREY, INK, MONO, NAVY, add_inline, set_cell_shade

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
HIST = os.path.join(ROOT, "docs", "env_v2", "history")
MD = os.path.join(HIST, "SYNTHETIC_ENVIRONMENT_HISTORY.md")
DOCX = os.path.join(HIST, "SYNTHETIC_ENVIRONMENT_HISTORY.docx")
PDF = os.path.join(HIST, "SYNTHETIC_ENVIRONMENT_HISTORY.pdf")

TEXT_WIDTH_CM = 17.6            # A4 (21.0 cm) minus two 1.7 cm margins
CAPTION = RGBColor(0x4B, 0x55, 0x63)
FONT = "Times New Roman"


def _force_font(doc):
    """Every run in the document, including tables, headers and footers, in the one typeface."""
    def runs_of(paragraphs):
        for p in paragraphs:
            for r in p.runs:
                yield r
    pars = list(doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                pars.extend(cell.paragraphs)
    for s in doc.sections:
        pars.extend(s.header.paragraphs); pars.extend(s.footer.paragraphs)
    for r in runs_of(pars):
        r.font.name = FONT
        rPr = r._element.get_or_add_rPr()
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = OxmlElement("w:rFonts"); rPr.append(rFonts)
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rFonts.set(qn(attr), FONT)


# ------------------------------------------------------------------------------------------------ small helpers
def _strip_md(text: str) -> str:
    """Headings and titles are plain text: drop inline markdown markers."""
    return re.sub(r"[`*]", "", text)


BOLD = re.compile(r"(\*\*.+?\*\*)")


def add_rich(par, text: str, size=10.5, color=INK, base_bold=False):
    """add_inline, but a bold span may itself contain inline code (the markdown tool's regex swallows it)."""
    for part in BOLD.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            add_inline(par, part[2:-2], size=size, color=color, base_bold=True)
        else:
            add_inline(par, part, size=size, color=color, base_bold=base_bold)


def add_heading(doc, text: str, level: int):
    """A Word heading (built-in style, so the table of contents sees it), coloured like the markdown tool's."""
    sizes = {1: 20, 2: 15, 3: 12.5, 4: 11}
    h = doc.add_heading(_strip_md(text), level=min(level, 4))
    for run in h.runs:
        run.font.color.rgb = NAVY; run.font.size = Pt(sizes.get(level, 11)); run.font.name = FONT; run.bold = True
    h.paragraph_format.space_before = Pt(16 if level <= 2 else 10); h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.keep_with_next = True
    if level == 2:
        pPr = h._p.get_or_add_pPr(); pbdr = OxmlElement("w:pBdr")
        b = OxmlElement("w:bottom"); b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "8"); b.set(qn("w:space"), "2"); b.set(qn("w:color"), "2A78D6")
        pbdr.append(b); pPr.append(pbdr)
    return h


def add_toc(doc):
    """A Word TOC field over heading levels 1 to 3; populated by Word when the fields are updated."""
    p = doc.add_paragraph()
    # the paragraph that keeps the field's end marker survives the update as an empty line after the last entry;
    # made one point tall so that it cannot spill onto a page of its own
    p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = Pt(1)
    run = p.add_run(); run.font.size = Pt(1)
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate")
    txt = OxmlElement("w:t"); txt.text = "The table of contents is filled in when the document's fields are updated."
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    for el in (begin, instr, sep, txt, end):
        run._r.append(el)
    return p


def _set_column_widths(t, cells, ncol):
    """Column widths in proportion to the text each column carries (the mean of its three longest cells, floored and
    capped), so that short columns stay narrow and long text gets the room justified lines need."""
    def longest_word(s):
        # capitals and digits are wider than lower-case letters in Times New Roman; weight them so that a status word
        # such as WITHDRAWN still fits on one line
        return max((sum(1.45 if c.isupper() else 1.0 for c in w) for w in s.split()), default=0)
    mins, wants = [], []
    for j in range(ncol):
        col = [(r[j] if j < len(r) else "") for r in cells]
        lw = min(25, max(longest_word(c) for c in col))
        lens = sorted((len(c) for c in col), reverse=True)[:3]
        mins.append(0.15 * lw + 0.5)                       # cm: the longest word on one line, plus the cell margins
        wants.append(0.14 * min(70, sum(lens) / len(lens)) + 0.5)
    total_min = sum(mins)
    if total_min >= TEXT_WIDTH_CM:
        widths = [TEXT_WIDTH_CM * m / total_min for m in mins]
    else:
        extra = [max(0.0, w - m) for w, m in zip(wants, mins)]
        pool, s = TEXT_WIDTH_CM - total_min, sum(extra)
        widths = [m + (pool * e / s if s > 0 else pool / ncol) for m, e in zip(mins, extra)]
    layout = OxmlElement("w:tblLayout"); layout.set(qn("w:type"), "fixed"); t._tbl.tblPr.append(layout)
    for j, w in enumerate(widths):
        t.columns[j].width = Cm(w)
        for i in range(len(cells)):
            t.cell(i, j).width = Cm(w)


def add_table(doc, rows):
    """A pipe table as a Word table: header shaded and repeated on every page, body at 8.5 pt (8 pt when wide)."""
    cells = [[c.strip() for c in re.split(r"(?<!\\)\|", r.strip().strip("|"))] for r in rows]
    cells = [r for r in cells if not all(re.fullmatch(r":?-{2,}:?", c or "---") for c in r)]
    if not cells:
        return
    ncol = max(len(r) for r in cells)
    size = 8.5 if ncol <= 5 else 8.0
    t = doc.add_table(rows=len(cells), cols=ncol)
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    _set_column_widths(t, cells, ncol)
    for i, r in enumerate(cells):
        for j in range(ncol):
            txt = r[j] if j < len(r) else ""
            cell = t.cell(i, j); cell.text = ""
            p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(1); p.paragraph_format.space_before = Pt(1)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_rich(p, txt.replace("\\|", "|"), size=size, base_bold=(i == 0), color=NAVY if i == 0 else INK)
            if i == 0:
                set_cell_shade(cell, "EEF2F8")
    # repeat the header row on every page the table spans
    trPr = t.rows[0]._tr.get_or_add_trPr()
    hdr = OxmlElement("w:tblHeader"); hdr.set(qn("w:val"), "true"); trPr.append(hdr)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_figure(doc, path: str, caption: str):
    """The figure at text width, centred, kept with its caption; the caption italic and grey."""
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True; p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(path, width=Cm(TEXT_WIDTH_CM))
    c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_inline(c, _unwrap_italic(caption), size=9, color=CAPTION)
    for run in c.runs:
        if run.font.name != MONO:
            run.italic = True
    c.paragraph_format.space_after = Pt(12)


def _unwrap_italic(text: str) -> str:
    """A whole paragraph wrapped in single asterisks: drop the wrapper so the inline code inside it still renders."""
    t = text.strip()
    if len(t) > 2 and t.startswith("*") and t.endswith("*") and not t.startswith("**"):
        return t[1:-1]
    return t


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _continuation(lines, j):
    return (j < len(lines) and lines[j].strip() and not lines[j].startswith(("#", "|", "```", "![")) and
            not re.match(r"^\s*([-*]|\d+\.)\s", lines[j]) and lines[j].strip() != "---")


# ------------------------------------------------------------------------------------------------ the conversion
def build_docx(md: str = MD, out: str = DOCX) -> str:
    lines = open(md, encoding="utf-8").read().splitlines()
    doc = Document()
    s = doc.sections[0]
    s.orientation = WD_ORIENT.PORTRAIT; s.page_width = Cm(21.0); s.page_height = Cm(29.7)
    s.left_margin = s.right_margin = Cm(1.7); s.top_margin = Cm(1.8); s.bottom_margin = Cm(1.8)
    base = doc.styles["Normal"]; base.font.name = FONT; base.font.size = Pt(11)
    base.paragraph_format.space_after = Pt(6); base.paragraph_format.line_spacing = 1.1
    base.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for name in ("List Bullet", "List Number"):
        doc.styles[name].font.name = FONT; doc.styles[name].font.size = Pt(11)
        doc.styles[name].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    i = 0
    title_done = False
    break_before_next = False
    skip_contents = False
    while i < len(lines):
        ln = lines[i].rstrip()
        if not ln.strip():
            i += 1; continue

        # the H1 becomes the title page: title, the status line under it, then the table of contents
        m = re.match(r"^#\s+(.*)$", ln)
        if m and not title_done:
            title = _strip_md(m.group(1))
            doc.core_properties.title = title; doc.core_properties.subject = "FinPersona synthetic market environment"
            tp = doc.add_paragraph(); tp.paragraph_format.space_before = Pt(120); tp.paragraph_format.space_after = Pt(18)
            tp.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = tp.add_run(title); r.bold = True; r.font.size = Pt(26); r.font.color.rgb = NAVY; r.font.name = FONT
            # the italic status paragraph that follows the title in the markdown
            j = i + 1
            while j < len(lines) and not lines[j].strip():
                j += 1
            if j < len(lines) and lines[j].startswith("*"):
                text = lines[j].strip(); k = j + 1
                while _continuation(lines, k):
                    text += " " + lines[k].strip(); k += 1
                sp = doc.add_paragraph(); add_inline(sp, _unwrap_italic(text), size=10.5, color=GREY)
                for run in sp.runs:
                    if run.font.name != MONO:
                        run.italic = True
                j = k
            page_break(doc)
            th = doc.add_paragraph(); tr = th.add_run("Table of contents")      # not a Heading style: keeps itself out of the TOC
            tr.bold = True; tr.font.size = Pt(20); tr.font.color.rgb = NAVY; tr.font.name = FONT
            th.paragraph_format.space_after = Pt(10)
            add_toc(doc)
            break_before_next = True                 # the first heading opens a new page without a stray blank page
            title_done = True; i = j; continue

        m = re.match(r"^(#{1,4})\s+(.*)$", ln)
        if m:
            level = len(m.group(1)); text = m.group(2)
            # the markdown's own contents list is replaced by the Word table of contents
            if level == 2 and text.strip().lower() == "contents":
                skip_contents = True; i += 1; continue
            skip_contents = False
            h = add_heading(doc, text, level)
            if break_before_next:
                h.paragraph_format.page_break_before = True; break_before_next = False
            i += 1; continue
        if skip_contents:
            i += 1; continue
        if ln.strip() == "---":
            i += 1; continue                                  # section rules: the headings carry their own border
        m = re.match(r"^!\[(.*?)\]\((.*?)\)\s*$", ln)
        if m:
            path = os.path.join(os.path.dirname(md), m.group(2))
            j = i + 1
            while j < len(lines) and not lines[j].strip():
                j += 1
            caption = ""
            if j < len(lines) and lines[j].startswith("*"):
                caption = lines[j].strip(); k = j + 1
                while _continuation(lines, k):
                    caption += " " + lines[k].strip(); k += 1
                j = k
            else:
                j = i + 1
            add_figure(doc, path, caption); i = j; continue
        if ln.startswith("```"):
            j = i + 1; block = []
            while j < len(lines) and not lines[j].startswith("```"):
                block.append(lines[j]); j += 1
            p = doc.add_paragraph(); run = p.add_run("\n".join(block)); run.font.name = MONO; run.font.size = Pt(8)
            run._element.rPr.rFonts.set(qn("w:eastAsia"), MONO); p.paragraph_format.space_after = Pt(6)
            i = j + 1; continue
        if ln.lstrip().startswith("|"):
            j = i; rows = []
            while j < len(lines) and lines[j].lstrip().startswith("|"):
                rows.append(lines[j]); j += 1
            add_table(doc, rows); i = j; continue
        m = re.match(r"^(\s*)[-*]\s+(.*)$", ln)
        if m:
            text = m.group(2); j = i + 1
            while _continuation(lines, j):
                text += " " + lines[j].strip(); j += 1
            p = doc.add_paragraph(style="List Bullet"); add_rich(p, text); p.paragraph_format.space_after = Pt(3)
            i = j; continue
        m = re.match(r"^(\d+)\.\s+(.*)$", ln)
        if m:
            text = m.group(2); j = i + 1
            while _continuation(lines, j):
                text += " " + lines[j].strip(); j += 1
            p = doc.add_paragraph(style="List Number"); add_rich(p, text); p.paragraph_format.space_after = Pt(3)
            i = j; continue
        text = ln.strip(); j = i + 1
        while _continuation(lines, j):
            text += " " + lines[j].strip(); j += 1
        p = doc.add_paragraph(); add_rich(p, text); i = j

    # page numbers in the footer; no running header
    for sec in doc.sections:
        fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = fp.add_run(); fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE"); run._r.append(fld)
        run.font.size = Pt(9); run.font.color.rgb = GREY
    _force_font(doc)
    doc.save(out)
    return out


def export_pdf(docx: str = DOCX, pdf: str = PDF) -> str:
    """Open the .docx in Word, update the fields and the table of contents, save it, export the PDF, close Word."""
    import win32com.client  # noqa: WPS433  (Windows only)
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False; word.DisplayAlerts = 0
    try:
        d = word.Documents.Open(os.path.abspath(docx), ReadOnly=False, AddToRecentFiles=False)
        # the built-in TOC styles (wdStyleTOC1..3) in the document's typeface, compact enough that the contents end
        # well clear of a page bottom
        for sid, size in ((-20, 11), (-21, 10.5), (-22, 10)):
            st = d.Styles(sid); st.Font.Name = FONT; st.Font.Size = size
            st.ParagraphFormat.SpaceBefore = 0; st.ParagraphFormat.SpaceAfter = 2; st.ParagraphFormat.LineSpacingRule = 0
        for k in range(1, d.TablesOfContents.Count + 1):
            d.TablesOfContents(k).Update()
        d.Fields.Update()
        d.Save()
        d.ExportAsFixedFormat(OutputFileName=os.path.abspath(pdf), ExportFormat=17, OpenAfterExport=False,
                              OptimizeFor=0, CreateBookmarks=1, DocStructureTags=True)
        pages = d.ComputeStatistics(2)   # wdStatisticPages
        d.Close(False)
    finally:
        word.Quit()
    time.sleep(0.5)
    return f"{pdf} ({pages} pages)"


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--md", default=MD); ap.add_argument("--docx", default=DOCX); ap.add_argument("--pdf", default=PDF)
    ap.add_argument("--no-pdf", action="store_true")
    a = ap.parse_args(argv)
    print("docx:", build_docx(a.md, a.docx))
    if not a.no_pdf:
        print("pdf:", export_pdf(a.docx, a.pdf))


if __name__ == "__main__":
    sys.exit(main())
